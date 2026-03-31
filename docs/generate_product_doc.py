#!/usr/bin/env python3
"""生成 BTCommunity 管理后台产品文档（Word 格式）"""

from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
import datetime


def set_cell_bg(cell, color_hex):
    """设置单元格背景色"""
    shading = cell._element.get_or_add_tcPr()
    shading_elem = shading.makeelement(qn("w:shd"), {
        qn("w:fill"): color_hex,
        qn("w:val"): "clear",
    })
    shading.append(shading_elem)


def add_table_row(table, cells_data, bold=False, header=False):
    """添加表格行"""
    row = table.add_row()
    for i, text in enumerate(cells_data):
        cell = row.cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        run = p.add_run(str(text))
        run.font.size = Pt(9)
        run.font.name = "微软雅黑"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        if bold or header:
            run.bold = True
        if header:
            set_cell_bg(cell, "2B579A")
            run.font.color.rgb = RGBColor(255, 255, 255)
    return row


def create_styled_table(doc, headers, rows):
    """创建带样式的表格"""
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"

    # 表头
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        run = p.add_run(h)
        run.bold = True
        run.font.size = Pt(9)
        run.font.name = "微软雅黑"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        run.font.color.rgb = RGBColor(255, 255, 255)
        set_cell_bg(cell, "2B579A")

    # 数据行
    for row_data in rows:
        add_table_row(table, row_data)

    return table


def add_heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.name = "微软雅黑"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    return h


def add_para(doc, text, bold=False, indent=False):
    p = doc.add_paragraph()
    if indent:
        p.paragraph_format.left_indent = Cm(0.5)
    run = p.add_run(text)
    run.font.size = Pt(10.5)
    run.font.name = "微软雅黑"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run.bold = bold
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet")
    p.clear()
    run = p.add_run(text)
    run.font.size = Pt(10)
    run.font.name = "微软雅黑"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    if level > 0:
        p.paragraph_format.left_indent = Cm(1.0 * level)
    return p


def main():
    doc = Document()

    # 页面设置
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)

    # ========== 封面 ==========
    for _ in range(6):
        doc.add_paragraph()

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("BTCommunity 管理后台")
    run.font.size = Pt(28)
    run.bold = True
    run.font.name = "微软雅黑"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run.font.color.rgb = RGBColor(43, 87, 154)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("产品需求文档（PRD）")
    run.font.size = Pt(18)
    run.font.name = "微软雅黑"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run.font.color.rgb = RGBColor(100, 100, 100)

    doc.add_paragraph()
    doc.add_paragraph()

    # 文档信息表
    info_table = doc.add_table(rows=5, cols=2)
    info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    info_table.style = "Table Grid"
    info_data = [
        ("文档版本", "V1.0"),
        ("产品名称", "BTCommunity 教育管理后台"),
        ("文档状态", "初稿"),
        ("创建日期", datetime.date.today().strftime("%Y-%m-%d")),
        ("技术栈", "Vue 3 + TypeScript + Element Plus + Vite"),
    ]
    for i, (k, v) in enumerate(info_data):
        for j, text in enumerate([k, v]):
            cell = info_table.rows[i].cells[j]
            cell.text = ""
            p = cell.paragraphs[0]
            run = p.add_run(text)
            run.font.size = Pt(10)
            run.font.name = "微软雅黑"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
            if j == 0:
                run.bold = True
                set_cell_bg(cell, "F2F2F2")

    doc.add_page_break()

    # ========== 目录页 ==========
    add_heading(doc, "目录", level=1)
    toc_items = [
        "1. 产品概述",
        "2. 用户角色与权限",
        "3. 功能模块详细说明",
        "   3.1 登录与认证",
        "   3.2 仪表盘",
        "   3.3 权限管理",
        "   3.4 成员管理",
        "   3.5 课程管理",
        "   3.6 班级管理",
        "   3.7 社区管理",
        "   3.8 通知管理",
        "   3.9 反馈管理",
        "   3.10 文档管理",
        "   3.11 文件管理",
        "   3.12 组织管理",
        "   3.13 系统设置",
        "   3.14 系统工具",
        "   3.15 监控日志",
        "4. 非功能性需求",
        "5. 部署与环境",
    ]
    for item in toc_items:
        p = doc.add_paragraph()
        run = p.add_run(item)
        run.font.size = Pt(10.5)
        run.font.name = "微软雅黑"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")

    doc.add_page_break()

    # ========== 1. 产品概述 ==========
    add_heading(doc, "1. 产品概述", level=1)

    add_heading(doc, "1.1 产品定位", level=2)
    add_para(doc,
        "BTCommunity 管理后台是一套面向教育培训机构的综合管理系统，"
        "为管理员提供对课程、班级、学员、教师、社区、通知、文档等核心业务的全生命周期管理能力。"
        "系统采用 RBAC（基于角色的访问控制）模型，支持多角色、多组织的灵活权限配置。"
    )

    add_heading(doc, "1.2 产品目标", level=2)
    goals = [
        "统一管理课程模板与班级课程排课，提升教务效率",
        "实现学员、教师、高级教师、组织员工的分类管理",
        "通过社区模块促进培训、合作、内部沟通",
        "提供多渠道通知系统（站内信/推送），确保信息触达",
        "完善的文档与资料管理，支持按用户/班级授权访问",
        "基于 RBAC 的细粒度权限控制，保障系统安全",
    ]
    for g in goals:
        add_bullet(doc, g)

    add_heading(doc, "1.3 技术架构", level=2)
    create_styled_table(doc,
        ["层级", "技术选型", "说明"],
        [
            ["前端框架", "Vue 3.5 + TypeScript 5.8", "Composition API + 严格类型"],
            ["UI 组件库", "Element Plus 2.10", "按需自动导入"],
            ["数据表格", "VXE-Table 4.6", "企业级高性能表格"],
            ["状态管理", "Pinia 3.0", "支持 setup/outside 双模式"],
            ["路由", "Vue Router 4.5", "支持动态路由加载"],
            ["HTTP 客户端", "Axios 1.10", "请求/响应拦截、超时重试"],
            ["CSS 方案", "UnoCSS + SCSS", "原子化 CSS"],
            ["构建工具", "Vite 7.0", "快速 HMR + 代码分割"],
            ["部署", "GitHub Actions + Self-hosted Runner", "CI/CD 自动部署到阿里云"],
        ]
    )

    doc.add_page_break()

    # ========== 2. 用户角色 ==========
    add_heading(doc, "2. 用户角色与权限", level=1)

    add_heading(doc, "2.1 角色定义", level=2)
    create_styled_table(doc,
        ["角色", "标识", "说明", "典型权限"],
        [
            ["超级管理员", "super_admin", "系统最高权限，管理所有模块", "全部功能"],
            ["管理员", "admin", "日常运营管理", "成员/课程/班级/通知管理"],
            ["编辑", "editor", "内容编辑与课程管理", "课程/文档编辑"],
            ["教师", "teacher", "课程教学相关操作", "查看所属班级/课程"],
            ["高级教师", "senior_teacher", "教学督导与班级管理", "跨班级查看/教学反馈"],
            ["学生", "student", "学习相关操作", "查看所属班级/课程"],
            ["班长", "class_monitor", "班级内部管理", "班级日常管理协助"],
            ["组织员工", "org_staff", "组织内部操作", "所属组织范围内操作"],
        ]
    )

    add_heading(doc, "2.2 权限模型", level=2)
    add_para(doc,
        "系统采用 RBAC + Casbin 权限模型：每个角色关联一组菜单权限（控制页面可见性）"
        "和 API 权限（控制接口访问）。管理员可在「角色管理」中动态配置。"
    )
    add_bullet(doc, "菜单权限：控制侧边栏菜单项的显示/隐藏，支持树形结构勾选")
    add_bullet(doc, "API 权限：控制后端接口的调用权限，按 API 分组管理")
    add_bullet(doc, "动态路由：前端路由从后端加载，确保无权限页面不可访问")

    doc.add_page_break()

    # ========== 3. 功能模块 ==========
    add_heading(doc, "3. 功能模块详细说明", level=1)

    # 3.1 登录
    add_heading(doc, "3.1 登录与认证", level=2)
    add_para(doc, "模块路径：/login", bold=True)
    add_para(doc, "功能描述：", bold=True)
    add_para(doc, "提供系统登录入口，支持邮箱/用户名 + 密码 + 图形验证码登录。登录成功后获取 JWT Token，存储于 localStorage，后续请求通过 x-token 请求头携带。")

    add_para(doc, "功能清单：", bold=True)
    create_styled_table(doc,
        ["功能点", "说明", "验证规则"],
        [
            ["邮箱登录", "输入注册邮箱和密码", "邮箱格式校验"],
            ["密码输入", "密码输入框", "6-20 位字符"],
            ["图形验证码", "后端生成验证码图片，用户输入后校验", "4 位字符，支持点击刷新"],
            ["账号状态检查", "登录时检查账号是否激活/禁用", "未激活或禁用账号提示相应信息"],
            ["Token 管理", "JWT Token 自动过期处理", "401 状态码自动登出并跳转登录页"],
        ]
    )

    add_para(doc, "业务流程：", bold=True)
    steps = [
        "1. 用户打开登录页，系统自动请求获取图形验证码",
        "2. 用户输入邮箱、密码、验证码，点击登录",
        "3. 后端校验凭据，返回 JWT Token",
        "4. 前端存储 Token，请求用户信息和动态菜单",
        "5. 根据角色跳转到对应仪表盘页面",
    ]
    for s in steps:
        add_bullet(doc, s)

    # 3.2 仪表盘
    add_heading(doc, "3.2 仪表盘", level=2)
    add_para(doc, "模块路径：/dashboard", bold=True)
    add_para(doc, "根据登录用户角色展示不同的仪表盘视图。super_admin 角色展示管理员视图，editor 角色展示编辑视图。作为系统首页，登录后默认跳转此页面。")

    # 3.3 权限管理
    add_heading(doc, "3.3 权限管理", level=2)
    add_para(doc, "模块路径：/authority", bold=True)
    add_para(doc, "包含用户管理、角色管理、菜单管理、API 管理四个子模块，构成完整的 RBAC 权限体系。")

    # 3.3.1 用户管理
    add_heading(doc, "3.3.1 用户管理", level=3)
    add_para(doc, "路径：/authority/user", bold=True)
    create_styled_table(doc,
        ["功能点", "操作", "说明"],
        [
            ["用户列表", "查看", "分页展示系统用户，支持按邮箱搜索"],
            ["新增用户", "创建", "填写邮箱、昵称、密码、角色"],
            ["编辑用户", "更新", "修改昵称、邮箱、性别、角色分配"],
            ["删除用户", "删除", "删除指定用户（需确认）"],
            ["修改密码", "更新", "管理员修改用户密码"],
            ["启用/禁用", "切换", "切换用户激活状态"],
            ["高级教师标记", "切换", "标记/取消用户的高级教师身份"],
        ]
    )

    add_para(doc, "列表字段：", bold=True)
    add_para(doc, "邮箱 | 昵称 | 激活状态 | 角色列表 | 创建时间")

    # 3.3.2 角色管理
    add_heading(doc, "3.3.2 角色管理", level=3)
    add_para(doc, "路径：/authority/role", bold=True)
    create_styled_table(doc,
        ["功能点", "操作", "说明"],
        [
            ["角色列表", "查看", "展示所有角色及基本信息"],
            ["新增角色", "创建", "输入角色名称和标识"],
            ["编辑角色", "更新", "修改角色名称和标识"],
            ["删除角色", "删除", "删除角色（需确认）"],
            ["菜单权限配置", "分配", "树形勾选该角色可访问的菜单"],
            ["API 权限配置", "分配", "勾选该角色可调用的 API 接口"],
        ]
    )

    # 3.3.3 菜单管理
    add_heading(doc, "3.3.3 菜单管理", level=3)
    add_para(doc, "路径：/authority/menu", bold=True)
    add_para(doc, "管理系统侧边栏菜单的树形结构。支持设置菜单名称、路由路径、组件路径、排序、图标（Element Plus 图标或自定义 SVG）、是否隐藏、是否缓存等属性。支持无限层级的父子菜单嵌套。")

    # 3.3.4 API 管理
    add_heading(doc, "3.3.4 API 管理", level=3)
    add_para(doc, "路径：/authority/api", bold=True)
    add_para(doc, "管理系统所有后端 API 接口的元数据，包括接口路径、请求方法、API 分组、描述。用于角色的 API 权限分配。")

    doc.add_page_break()

    # 3.4 成员管理
    add_heading(doc, "3.4 成员管理", level=2)
    add_para(doc, "模块路径：/member", bold=True)
    add_para(doc, "管理系统中四类业务用户：学生、教师、高级教师、组织员工。每类成员独立管理，共享相似的操作模式。")

    member_types = [
        ("3.4.1 学生管理", "/member/student", "student", [
            ["学生列表", "查看", "分页展示，支持按邮箱/昵称/组织筛选"],
            ["新增学生", "创建", "填写邮箱、昵称、性别、组织、密码"],
            ["编辑学生", "更新", "修改学生基本信息和组织归属"],
            ["删除学生", "删除", "删除学生账号"],
            ["重置密码", "更新", "管理员重置学生登录密码"],
            ["启用/禁用", "切换", "切换学生账号激活状态"],
        ]),
        ("3.4.2 教师管理", "/member/teacher", "teacher", [
            ["教师列表", "查看", "分页展示，支持按邮箱/昵称/组织筛选"],
            ["新增教师", "创建", "填写邮箱、昵称、性别、组织、地点、密码"],
            ["编辑教师", "更新", "修改教师基本信息"],
            ["删除教师", "删除", "删除教师账号"],
            ["重置密码", "更新", "管理员重置教师登录密码"],
            ["启用/禁用", "切换", "切换教师账号激活状态"],
        ]),
        ("3.4.3 高级教师管理", "/member/seniorTeacher", "senior_teacher", [
            ["高级教师列表", "查看", "分页展示，支持按邮箱/昵称/组织筛选"],
            ["新增高级教师", "创建", "填写邮箱、昵称、性别、组织、密码"],
            ["编辑高级教师", "更新", "修改高级教师基本信息"],
            ["删除高级教师", "删除", "删除高级教师账号"],
            ["重置密码", "更新", "管理员重置登录密码"],
            ["启用/禁用", "切换", "切换账号激活状态"],
        ]),
        ("3.4.4 组织员工管理", "/member/orgStaff", "org_staff", [
            ["员工列表", "查看", "分页展示，支持按邮箱/昵称/组织筛选"],
            ["新增员工", "创建", "填写邮箱、昵称、性别、组织、地点、密码"],
            ["编辑员工", "更新", "修改员工基本信息"],
            ["删除员工", "删除", "删除员工账号"],
            ["重置密码", "更新", "管理员重置登录密码"],
            ["启用/禁用", "切换", "切换账号激活状态"],
        ]),
    ]

    for title, path, _, features in member_types:
        add_heading(doc, title, level=3)
        add_para(doc, f"路径：{path}", bold=True)
        create_styled_table(doc,
            ["功能点", "操作", "说明"],
            features
        )
        add_para(doc, "列表字段：邮箱 | 昵称 | 性别 | 所属组织 | 激活状态 | 创建时间 | 最后登录时间")

    doc.add_page_break()

    # 3.5 课程管理
    add_heading(doc, "3.5 课程管理", level=2)
    add_para(doc, "模块路径：/course", bold=True)
    add_para(doc, "课程管理是系统核心业务模块，支持课程模板的创建、编辑、分类管理和课时排课。课程采用「课程 → 章节 → 课时」三级结构。")

    add_heading(doc, "3.5.1 课程列表", level=3)
    create_styled_table(doc,
        ["功能点", "操作", "说明"],
        [
            ["课程列表", "查看", "分页展示课程模板，支持按分类筛选"],
            ["新增课程", "创建", "填写课程名称、描述、分类"],
            ["编辑课程", "更新", "修改课程基本信息"],
            ["删除课程", "删除", "删除课程模板"],
            ["排课", "操作", "进入课程目录管理，配置章节和课时"],
        ]
    )
    add_para(doc, "列表字段：课程名称 | 分类 | 创建者 | 创建时间 | 课时数 | 状态")

    add_heading(doc, "3.5.2 课程分类管理", level=3)
    add_para(doc, "通过弹窗管理课程分类，支持新增、编辑、删除分类。分类用于课程列表的筛选和归类。")

    add_heading(doc, "3.5.3 章节管理", level=3)
    add_para(doc, "在课程详情中管理章节，支持新增、编辑、删除、排序。章节是课时的容器，用于组织课程内容结构。")

    add_heading(doc, "3.5.4 课时管理", level=3)
    create_styled_table(doc,
        ["功能点", "操作", "说明"],
        [
            ["新增课时", "创建", "填写课时名称、类型、所属章节"],
            ["编辑课时", "更新", "修改课时内容、时间、教师等"],
            ["删除课时", "删除", "删除指定课时"],
            ["视频上传", "上传", "支持大文件分片上传（Chunk Upload）"],
            ["富文本内容", "编辑", "使用 Quill 编辑器编写课时内容"],
            ["附件管理", "上传", "支持多种格式附件上传"],
        ]
    )
    add_para(doc, "课时字段：课时名称 | 课时类型 | 所属章节 | 上课时间 | 时长 | 授课教师 | 视频 | 附件 | 讨论模式 | 直播链接")

    doc.add_page_break()

    # 3.6 班级管理
    add_heading(doc, "3.6 班级管理", level=2)
    add_para(doc, "模块路径：/class", bold=True)
    add_para(doc, "班级是教学的实际组织单元，将课程模板实例化后分配给具体的教师和学生。")

    add_heading(doc, "3.6.1 班级列表", level=3)
    create_styled_table(doc,
        ["功能点", "操作", "说明"],
        [
            ["班级列表", "查看", "分页展示，支持按名称/教师/学生搜索"],
            ["新增班级", "创建", "填写名称、时间范围、备注，选择教师和学生"],
            ["班级详情", "查看", "进入班级详情页管理"],
            ["删除班级", "删除", "删除班级"],
        ]
    )
    add_para(doc, "列表字段：班级名称 | 教师数 | 高级教师数 | 学生数 | 课时进度")

    add_heading(doc, "3.6.2 班级详情（Tab 页）", level=3)
    create_styled_table(doc,
        ["Tab 页", "功能", "说明"],
        [
            ["基本信息", "编辑班级名称、时间、组织、备注", "直接在详情页编辑"],
            ["课程管理", "导入课程模板、管理班级课程", "从课程模板导入，管理章节和课时"],
            ["资料管理", "查看班级已授权的文档资料", "关联文档管理模块的授权"],
            ["学生管理", "添加/移除班级学生", "从成员库选择学生加入班级"],
            ["教师管理", "添加/移除班级教师和高级教师", "从成员库选择教师加入班级"],
        ]
    )

    doc.add_page_break()

    # 3.7 社区管理
    add_heading(doc, "3.7 社区管理", level=2)
    add_para(doc, "模块路径：/community", bold=True)
    add_para(doc, "社区模块提供即时通讯（IM）能力，支持三种预设社区类型，每个社区包含成员管理和会话管理。")

    add_heading(doc, "3.7.1 社区类型", level=3)
    create_styled_table(doc,
        ["类型", "标识", "说明"],
        [
            ["培训社区", "Training (1)", "面向培训班的沟通社区"],
            ["合作社区", "Cooperation (2)", "跨组织合作交流社区"],
            ["员工社区", "Employee (3)", "内部员工沟通社区"],
        ]
    )

    add_heading(doc, "3.7.2 社区列表", level=3)
    add_para(doc, "以卡片形式展示三种社区类型，每种类型显示统计数据：社区数量、独立成员数、会话数。点击卡片进入该类型的社区管理。")

    add_heading(doc, "3.7.3 社区管理详情", level=3)
    create_styled_table(doc,
        ["功能点", "操作", "说明"],
        [
            ["成员管理", "增/删/改/查", "管理社区成员，支持按角色分组查看"],
            ["会话管理", "增/删/改/查", "管理社区内的群聊/公告群"],
            ["会话类型", "配置", "普通群聊 (1) / 公告群 (2)"],
            ["会话同步", "操作", "同步会话到 IM 服务（XMPP）"],
            ["会话归档", "操作", "归档不再使用的会话"],
            ["成员候选人", "查询", "从系统用户中搜索可添加的成员"],
        ]
    )

    doc.add_page_break()

    # 3.8 通知管理
    add_heading(doc, "3.8 通知管理", level=2)
    add_para(doc, "模块路径：/notice", bold=True)
    add_para(doc, "多渠道通知系统，支持定时发送、目标筛选、送达统计。")

    add_heading(doc, "3.8.1 通知列表", level=3)
    add_para(doc, "按状态分 Tab 展示：草稿 | 定时 | 已发布 | 已取消 | 已过期。已发布 Tab 下再分「手动」和「自动」子标签。")
    add_para(doc, "列表字段：标题 | 类型 | 分类 | 状态 | 优先级 | 目标数 | 已发送数 | 已读数 | 创建时间")

    add_heading(doc, "3.8.2 创建通知", level=3)
    create_styled_table(doc,
        ["字段", "类型", "说明"],
        [
            ["通知类型", "选择", "系统通知 / 业务通知 / 告警通知"],
            ["分类", "选择", "课程安排、课程变更、作业创建等 14 种预设分类"],
            ["标题", "文本", "通知标题"],
            ["内容", "富文本", "Quill 编辑器输入通知正文"],
            ["优先级", "选择", "普通 (1) / 重要 (2) / 紧急 (3)"],
            ["目标类型", "选择", "用户 / 角色 / 组织 / 班级 / 社区 / 全部"],
            ["目标选择", "弹窗选择", "根据目标类型选择具体接收对象"],
            ["发送渠道", "多选", "从后端获取可用渠道列表"],
            ["定时发送", "开关", "开启后选择发送时间和时区"],
            ["过期时间", "日期", "通知过期后不再展示"],
        ]
    )

    add_heading(doc, "3.8.3 通知操作", level=3)
    create_styled_table(doc,
        ["操作", "条件", "说明"],
        [
            ["发送", "草稿状态", "立即发送或设定定时"],
            ["编辑", "草稿/定时状态", "修改通知内容和目标"],
            ["取消", "定时状态", "取消定时发送"],
            ["删除", "所有状态", "删除通知记录"],
            ["查看详情", "所有状态", "查看通知内容和送达统计"],
        ]
    )

    doc.add_page_break()

    # 3.9 反馈管理
    add_heading(doc, "3.9 反馈管理", level=2)
    add_para(doc, "模块路径：/feedback", bold=True)

    add_heading(doc, "3.9.1 用户反馈", level=3)
    add_para(doc, "路径：/feedback/list", bold=True)
    create_styled_table(doc,
        ["功能点", "操作", "说明"],
        [
            ["反馈列表", "查看", "分页展示用户反馈，支持按类型筛选"],
            ["查看详情", "查看", "查看反馈内容、截图、诊断日志"],
            ["删除反馈", "删除", "删除指定反馈记录"],
        ]
    )
    add_para(doc, "反馈类型：Bug (1) | 建议 (2) | 其他 (3)")
    add_para(doc, "列表字段：反馈类型 | 提交人 | 创建时间 | 描述")

    add_heading(doc, "3.9.2 教学反馈", level=3)
    add_para(doc, "路径：/feedback/teaching", bold=True)
    add_para(doc, "独立的教学质量反馈系统，记录教学相关的意见和建议。支持查看列表、详情和删除。")

    # 3.10 文档管理
    add_heading(doc, "3.10 文档管理", level=2)
    add_para(doc, "模块路径：/document", bold=True)
    add_para(doc, "树形结构的文档资料管理系统，支持文件夹嵌套、文件上传、按用户/班级授权访问。")

    create_styled_table(doc,
        ["功能点", "操作", "说明"],
        [
            ["文件夹管理", "创建/编辑/删除", "树形目录结构管理"],
            ["文件上传", "上传", "支持 xlsx/csv/txt/mp3/docx/mp4/pdf/jpg/png 等格式"],
            ["云端上传", "上传", "支持直传到云存储（OSS）"],
            ["重复检测", "校验", "上传前检查同目录下是否有同名文件"],
            ["搜索", "查询", "按关键词搜索文档"],
            ["用户授权", "分配", "将文档/文件夹授权给指定用户"],
            ["班级授权", "分配", "将文档/文件夹授权给指定班级"],
            ["批量取消授权", "操作", "批量撤销已有授权"],
        ]
    )

    # 3.11 文件管理
    add_heading(doc, "3.11 文件管理", level=2)
    add_para(doc, "模块路径：/fileM", bold=True)
    add_para(doc, "系统底层文件存储管理，展示所有已上传的文件。支持搜索、下载、删除操作。")
    add_para(doc, "列表字段：文件名 | 大小 | MD5 | 上传时间")
    add_para(doc, "上传能力：", bold=True)
    add_bullet(doc, "普通上传：小文件直接上传")
    add_bullet(doc, "分片上传（Chunk Upload）：大文件分片上传，支持断点续传、进度查询、取消上传")
    add_bullet(doc, "云端直传（Direct Upload）：客户端直传到 OSS，减轻服务器带宽")
    add_bullet(doc, "分段上传（Multipart Upload）：云端分段上传，支持进度跟踪")

    doc.add_page_break()

    # 3.12 组织管理
    add_heading(doc, "3.12 组织管理", level=2)
    add_para(doc, "模块路径：/organization", bold=True)
    create_styled_table(doc,
        ["功能点", "操作", "说明"],
        [
            ["组织列表", "查看", "分页展示，支持按名称搜索"],
            ["新增组织", "创建", "输入组织名称和备注"],
            ["编辑组织", "更新", "修改名称和备注"],
            ["删除组织", "删除", "删除组织"],
        ]
    )
    add_para(doc, "组织用于关联成员（学生、教师、员工），实现基于组织维度的人员管理和数据隔离。")

    # 3.13 系统设置
    add_heading(doc, "3.13 系统设置", level=2)
    add_para(doc, "模块路径：/sysSet", bold=True)

    add_heading(doc, "3.13.1 数据字典", level=3)
    add_para(doc, "管理系统中的枚举数据（如性别、状态等），采用左右分栏布局：左侧字典列表，右侧字典条目管理。")
    create_styled_table(doc,
        ["功能点", "操作", "说明"],
        [
            ["字典列表", "查看", "展示所有数据字典（中文名/英文名）"],
            ["新增字典", "创建", "输入中文名和英文标识"],
            ["编辑/删除字典", "操作", "修改或删除字典"],
            ["字典条目管理", "CRUD", "管理字典下的 label/value/排序"],
        ]
    )

    # 3.14 系统工具
    add_heading(doc, "3.14 系统工具", level=2)
    add_para(doc, "模块路径：/sysTool", bold=True)

    add_heading(doc, "3.14.1 定时任务管理", level=3)
    create_styled_table(doc,
        ["功能点", "操作", "说明"],
        [
            ["任务列表", "查看", "分页展示所有定时任务"],
            ["新增任务", "创建", "配置名称、Cron 表达式、执行方法、策略"],
            ["编辑任务", "更新", "修改任务配置"],
            ["删除任务", "删除", "单条或批量删除"],
            ["启用/禁用", "切换", "控制任务的运行状态"],
        ]
    )
    add_para(doc, "执行方法：ClearTable（清理数据表）| Shell（执行 Shell 命令）")
    add_para(doc, "执行策略：始终执行 | 仅失败时执行")

    # 3.15 监控日志
    add_heading(doc, "3.15 监控日志", level=2)
    add_para(doc, "模块路径：/monitor", bold=True)

    add_heading(doc, "3.15.1 操作日志", level=3)
    add_para(doc, "记录所有 API 调用的审计日志，支持多条件筛选和批量删除。")
    create_styled_table(doc,
        ["功能点", "操作", "说明"],
        [
            ["日志列表", "查看", "分页展示，支持按路径/方法/状态码筛选"],
            ["排序", "操作", "按响应时间升序/降序"],
            ["删除", "操作", "单条或批量删除日志"],
        ]
    )
    add_para(doc, "日志字段：IP | 请求方法 | 路径 | 状态码 | 操作用户 | 时间 | 响应时间 | UserAgent | 请求参数 | 响应数据")

    doc.add_page_break()

    # ========== 4. 非功能性需求 ==========
    add_heading(doc, "4. 非功能性需求", level=1)

    add_heading(doc, "4.1 响应式布局", level=2)
    add_para(doc, "系统支持三种布局模式（左侧栏、顶部栏、混合模式），自动检测设备类型（桌面/平板/移动端）并适配。布局偏好持久化存储于 localStorage。")

    add_heading(doc, "4.2 主题与个性化", level=2)
    create_styled_table(doc,
        ["配置项", "说明"],
        [
            ["布局模式", "左侧栏 / 顶部栏 / 混合模式"],
            ["暗色模式", "支持明暗主题切换"],
            ["标签栏", "显示/隐藏页面标签栏"],
            ["水印", "开启/关闭系统水印"],
            ["灰色模式", "全站灰度显示"],
            ["色弱模式", "色弱友好模式"],
            ["固定头部", "头部是否固定"],
        ]
    )

    add_heading(doc, "4.3 安全性", level=2)
    add_bullet(doc, "JWT Token 认证，401 自动登出")
    add_bullet(doc, "图形验证码防暴力破解")
    add_bullet(doc, "RBAC + Casbin 细粒度权限控制")
    add_bullet(doc, "前端动态路由，无权限页面不可访问")
    add_bullet(doc, "API 请求拦截器自动注入认证 Token")
    add_bullet(doc, "生产环境自动移除 console/debugger")

    add_heading(doc, "4.4 性能优化", level=2)
    add_bullet(doc, "代码分割：Vue/Element Plus/VXE-Table 独立 chunk")
    add_bullet(doc, "组件按需加载：路由懒加载 + 组件自动导入")
    add_bullet(doc, "大文件分片上传：支持断点续传")
    add_bullet(doc, "Gzip 压缩支持")
    add_bullet(doc, "静态资源长期缓存（immutable）")

    doc.add_page_break()

    # ========== 5. 部署与环境 ==========
    add_heading(doc, "5. 部署与环境", level=1)

    add_heading(doc, "5.1 环境配置", level=2)
    create_styled_table(doc,
        ["环境", "构建命令", "触发分支", "部署路径"],
        [
            ["开发（本地）", "pnpm dev", "-", "localhost:5173"],
            ["Staging", "pnpm build:staging", "dev", "/www/wwwroot/lms7799"],
            ["Production", "pnpm build", "release", "/var/www/btcommunity-web"],
        ]
    )

    add_heading(doc, "5.2 CI/CD 流程", level=2)
    create_styled_table(doc,
        ["工作流", "触发条件", "运行环境", "作用"],
        [
            ["ci.yml", "Push/PR 到 dev/main", "GitHub 托管 (ubuntu)", "ESLint + 类型检查 + 测试 + 构建验证"],
            ["deploy-staging.yml", "Push 到 dev", "Self-hosted Runner (Docker)", "构建并部署到 Staging"],
            ["deploy-production.yml", "Push 到 release", "Self-hosted Runner (Docker)", "构建并部署到 Production"],
        ]
    )

    add_heading(doc, "5.3 服务器要求", level=2)
    add_bullet(doc, "阿里云 ECS，安装 Docker + Nginx")
    add_bullet(doc, "Nginx 配置 SPA history 模式（try_files）")
    add_bullet(doc, "Nginx 配置 /api 反向代理到后端服务")
    add_bullet(doc, "GitHub Self-hosted Runner 容器化部署")
    add_bullet(doc, "部署目录通过 Docker Volume 挂载")

    # 保存
    output_path = "/Users/jianjieguo/Documents/GitHub/academic-web/docs/BTCommunity_产品文档_V1.0.docx"
    doc.save(output_path)
    print(f"文档已生成: {output_path}")


if __name__ == "__main__":
    main()